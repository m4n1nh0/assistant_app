"""Leitura de material didatico: limpeza, corte e erro honesto.

O gerador de quiz recebe esse texto no prompt. Texto picado por quebra de linha
de PDF produz pergunta sobre fragmento; texto grande demais volta cortado do
modelo e derruba o quiz no gerador por template. As duas coisas se resolvem
aqui, antes de chegar na IA.

Os arquivos de DOCX e PPTX sao montados de verdade em memoria, e nao simulados:
o que se quer verificar e justamente a leitura do formato real.
"""

from __future__ import annotations

import asyncio
import io

import pytest

from app.services import material_service as ms


class _Page:
    def __init__(self, text: str = "", erro: Exception | None = None):
        self._text = text
        self._erro = erro

    def extract_text(self) -> str:
        if self._erro:
            raise self._erro
        return self._text


def _texto_longo(prefixo: str = "") -> str:
    return prefixo + " ".join(
        f"A juncao interna retorna somente as linhas que atendem a condicao {i}."
        for i in range(1, 12)
    )


def _docx_bytes(paragrafos: list[str], tabela: list[list[str]] | None = None) -> bytes:
    from docx import Document

    documento = Document()
    for texto in paragrafos:
        documento.add_paragraph(texto)
    if tabela:
        grade = documento.add_table(rows=len(tabela), cols=len(tabela[0]))
        for i, linha in enumerate(tabela):
            for j, celula in enumerate(linha):
                grade.cell(i, j).text = celula

    buffer = io.BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


#: Texto impresso nas paginas sinteticas. Passa de MIN_CHARS com folga, para o
#: teste nao virar refem de uma linha que o OCR eventualmente descarte.
_LINHAS_AULA = [
    "Normalizacao de banco de dados",
    "A primeira forma normal elimina grupos repetitivos.",
    "A segunda forma normal trata da dependencia parcial.",
    "A terceira forma normal elimina dependencia transitiva.",
    "Chave primaria garante a identidade de cada linha.",
    "Integridade referencial liga a chave estrangeira a origem.",
]


def _pagina_impressa(linhas: list[str]):
    """Desenha texto numa imagem, como sairia de um escaner."""
    from PIL import Image, ImageDraw, ImageFont

    imagem = Image.new("RGB", (1240, 90 + 70 * len(linhas)), "white")
    desenho = ImageDraw.Draw(imagem)
    try:
        fonte = ImageFont.truetype("arial.ttf", 32)
    except Exception:  # pragma: no cover - depende das fontes da maquina
        fonte = ImageFont.load_default()

    y = 50
    for linha in linhas:
        desenho.text((50, y), linha, fill="black", font=fonte)
        y += 70
    return imagem


def _png_bytes(linhas: list[str]) -> bytes:
    buffer = io.BytesIO()
    _pagina_impressa(linhas).save(buffer, format="PNG")
    return buffer.getvalue()


def _pdf_digitalizado(paginas: list[list[str]]) -> bytes:
    """PDF feito de imagens: e o que um escaner produz, sem camada de texto."""
    imagens = [_pagina_impressa(linhas) for linhas in paginas]
    buffer = io.BytesIO()
    imagens[0].save(
        buffer,
        format="PDF",
        save_all=True,
        append_images=imagens[1:],
        resolution=150,
    )
    return buffer.getvalue()


def _pptx_bytes(slides: list[tuple[str, list[str], str]]) -> bytes:
    """Monta um deck: cada slide e (titulo, bullets, notas do apresentador)."""
    from pptx import Presentation

    deck = Presentation()
    layout = deck.slide_layouts[1]  # titulo e conteudo
    for titulo, bullets, notas in slides:
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = titulo
        corpo = slide.placeholders[1].text_frame
        corpo.text = bullets[0]
        for bullet in bullets[1:]:
            corpo.add_paragraph().text = bullet
        if notas:
            slide.notes_slide.notes_text_frame.text = notas

    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


# --- Limpeza ----------------------------------------------------------------


def test_junta_linha_quebrada_pela_largura_da_pagina():
    # PDF quebra por largura, nao por sentido: sem juntar, o modelo recebe meio
    # conceito por linha.
    limpo = ms.clean_text("A juncao de tabelas e uma\ndas operacoes do modelo\nrelacional.")

    assert limpo == "A juncao de tabelas e uma das operacoes do modelo relacional."


def test_junta_palavra_hifenizada_no_fim_da_linha():
    limpo = ms.clean_text("O produto carte-\nsiano combina tudo.")

    assert "cartesiano" in limpo


def test_preserva_a_divisao_entre_paragrafos():
    # E ela que separa um assunto do proximo; achatar tudo viraria um bloco so.
    limpo = ms.clean_text("Primeiro assunto\ncontinua aqui.\n\nSegundo assunto.")

    assert limpo == "Primeiro assunto continua aqui.\n\nSegundo assunto."


# --- Extracao ---------------------------------------------------------------


def test_pagina_ilegivel_nao_perde_o_material_inteiro():
    paginas = [
        _Page(_texto_longo("Pagina boa. ")),
        _Page(erro=ValueError("fonte corrompida")),
        _Page(_texto_longo("Outra pagina boa. ")),
    ]

    extraido = ms.from_pages(paginas)

    assert extraido.page_count == 2
    assert "Pagina boa" in extraido.text
    assert "Outra pagina boa" in extraido.text


def test_pdf_digitalizado_falha_com_motivo():
    # PDF de imagem extrai vazio. Gravar material em branco daria erro so la na
    # frente, na geracao do quiz, sem explicar a causa.
    with pytest.raises(ms.MaterialError, match="OCR"):
        ms.from_pages([_Page(""), _Page("   ")])


def test_material_grande_e_cortado_e_avisa(monkeypatch):
    monkeypatch.setattr(ms, "MAX_CHARS", 400)

    extraido = ms.from_pages([_Page(_texto_longo()), _Page(_texto_longo())])

    assert extraido.truncated is True
    assert extraido.char_count <= 400


def test_arquivo_vazio_nao_vira_material():
    with pytest.raises(ms.MaterialError):
        asyncio.run(ms.extract_pdf(b""))


def test_pdf_ilegivel_falha_sem_estourar_excecao_crua():
    with pytest.raises(ms.MaterialError, match="Nao consegui ler"):
        asyncio.run(ms.extract_pdf(b"isto nao e um PDF"))


# --- Formatos estruturados ---------------------------------------------------


def test_linha_com_sentido_proprio_nao_e_juntada():
    """A diferenca entre PDF e slide.

    No PDF a quebra e acidente de largura de pagina e precisa ser desfeita; num
    slide ela separa dois bullets, e junta-los entregaria ao gerador um
    paragrafo corrido no lugar de dois topicos.
    """
    bullets = "Chave primaria identifica a linha\nChave estrangeira liga tabelas"

    assert ms.clean_text(bullets, rewrap=True) == (
        "Chave primaria identifica a linha Chave estrangeira liga tabelas"
    )
    assert ms.clean_text(bullets, rewrap=False) == bullets


def test_docx_le_paragrafos_e_tabela():
    data = _docx_bytes(
        [_texto_longo("Conceito. "), _texto_longo("Aplicacao. ")],
        tabela=[["Operador", "Efeito"], ["INNER JOIN", "So o que casa"]],
    )

    extraido = asyncio.run(ms.extract(data, "apostila.docx"))

    assert extraido.source_type == "docx"
    assert "Conceito" in extraido.text and "Aplicacao" in extraido.text
    # A tabela vira linha legivel em vez de sumir: e nela que costuma estar o
    # resumo comparativo que rende boa pergunta.
    assert "INNER JOIN | So o que casa" in extraido.text
    # DOCX nao tem pagina antes de ser renderizado; contar paragrafos daria um
    # numero sem sentido na interface.
    assert extraido.page_count == 1


def test_docx_vazio_nao_manda_o_professor_procurar_OCR():
    # A dica de digitalizacao so vale para quem pode chegar como imagem.
    with pytest.raises(ms.MaterialError) as erro:
        asyncio.run(ms.extract(_docx_bytes(["", "  "]), "vazio.docx"))

    assert "OCR" not in str(erro.value)


def test_pptx_conta_um_bloco_por_slide():
    data = _pptx_bytes(
        [
            ("Juncoes", [_texto_longo("Primeiro slide. ")], ""),
            ("Agregacoes", [_texto_longo("Segundo slide. ")], ""),
        ]
    )

    extraido = asyncio.run(ms.extract(data, "aula03.pptx"))

    assert extraido.source_type == "pptx"
    assert extraido.page_count == 2


def test_pptx_aproveita_as_notas_do_apresentador():
    # O bullet resume; a nota traz a explicacao que vira a pergunta.
    data = _pptx_bytes(
        [
            (
                "Normalizacao",
                [_texto_longo("Bullet. ")],
                "A terceira forma normal elimina dependencia transitiva.",
            )
        ]
    )

    extraido = asyncio.run(ms.extract(data, "aula.pptx"))

    assert "dependencia transitiva" in extraido.text


def test_markdown_preserva_os_itens_da_lista():
    conteudo = f"# Banco de dados\n- Chave primaria\n- Chave estrangeira\n\n{_texto_longo()}"

    extraido = asyncio.run(ms.extract(conteudo.encode("utf-8"), "resumo.md"))

    assert extraido.source_type == "md"
    assert "- Chave primaria\n- Chave estrangeira" in extraido.text


def test_txt_salvo_no_windows_nao_e_recusado_por_um_acento():
    # latin-1 e o que sai do Bloco de Notas antigo; recusar perderia o material.
    data = _texto_longo("Introdução à normalização. ").encode("latin-1")

    extraido = asyncio.run(ms.extract(data, "notas.txt"))

    assert extraido.source_type == "txt"


def test_extensao_e_reconhecida_independente_da_caixa():
    assert ms.is_supported("APOSTILA.PDF")
    assert ms.extension_of("Aula 03.PPTX") == ".pptx"


def test_formato_nao_suportado_diz_o_que_e_aceito():
    # Sem a lista, o professor tenta de novo com outro arquivo qualquer.
    with pytest.raises(ms.MaterialError) as erro:
        asyncio.run(ms.extract(b"conteudo", "notas.xlsx"))

    assert ".xlsx" in str(erro.value)
    assert ".pptx" in str(erro.value)


# --- Material que chega como imagem (OCR) ------------------------------------


def test_pdf_digitalizado_confirma_que_nao_tem_camada_de_texto():
    """Ancora dos testes de OCR abaixo.

    Se um dia o `pypdf` passar a extrair texto desse arquivo, os testes
    seguintes continuariam verdes sem nunca exercitar o OCR.
    """
    from pypdf import PdfReader

    paginas = PdfReader(io.BytesIO(_pdf_digitalizado([_LINHAS_AULA]))).pages

    assert not (paginas[0].extract_text() or "").strip()


def test_ocr_desligado_recusa_o_digitalizado_com_o_motivo_de_sempre(monkeypatch):
    monkeypatch.setattr(ms.ocr_service, "is_available", lambda: False)

    with pytest.raises(ms.MaterialError, match="OCR"):
        asyncio.run(ms.extract(_pdf_digitalizado([_LINHAS_AULA]), "escaneado.pdf"))


def test_pdf_com_texto_nao_paga_o_custo_do_ocr(monkeypatch):
    # OCR custa cerca de um segundo por pagina. Rodar quando a camada de texto
    # ja resolveu seria desperdicio puro no upload.
    pronto = ms.ExtractedMaterial(
        text="conteudo ja extraido " * 20, page_count=2, truncated=False
    )
    monkeypatch.setattr(ms, "from_pages", lambda pages: pronto)

    def _nao_deveria_ser_chamado(*args, **kwargs):
        raise AssertionError("OCR acionado para PDF que ja tinha texto")

    monkeypatch.setattr(ms.ocr_service, "text_from_pdf", _nao_deveria_ser_chamado)

    assert ms.extract_pdf_sync(_pdf_digitalizado([_LINHAS_AULA])) is pronto


def test_imagem_sem_ocr_explica_que_o_recurso_esta_desligado(monkeypatch):
    monkeypatch.setattr(ms.ocr_service, "is_available", lambda: False)

    with pytest.raises(ms.MaterialError, match="desligado"):
        asyncio.run(ms.extract(b"qualquer coisa", "quadro.jpg"))


def test_imagem_ilegivel_recusa_com_motivo_que_o_professor_entende():
    with pytest.raises(ms.MaterialError, match="tremida"):
        asyncio.run(ms.extract(b"isto nao e uma imagem", "quadro.png"))


@pytest.mark.integration
def test_pdf_digitalizado_e_lido_pelo_ocr():
    pytest.importorskip("rapidocr")
    pytest.importorskip("pypdfium2")

    extraido = asyncio.run(
        ms.extract(_pdf_digitalizado([_LINHAS_AULA]), "apostila-escaneada.pdf")
    )

    # O formato registrado distingue o texto lido por OCR do texto nativo: ele
    # pode ter erro de reconhecimento, e o professor merece saber disso.
    assert extraido.source_type == "pdf-ocr"
    assert "transitiva" in extraido.text
    assert extraido.page_count == 1


@pytest.mark.integration
def test_teto_de_paginas_marca_o_material_como_cortado(monkeypatch):
    # Sem o teto, uma apostila digitalizada de duzentas paginas seguraria o
    # upload por minutos ate o cliente desistir.
    pytest.importorskip("rapidocr")
    monkeypatch.setattr(ms.ocr_service.settings, "ocr_max_pages", 1, raising=False)

    extraido = asyncio.run(
        ms.extract(_pdf_digitalizado([_LINHAS_AULA, _LINHAS_AULA]), "longa.pdf")
    )

    assert extraido.page_count == 1
    assert extraido.truncated is True


@pytest.mark.integration
def test_foto_do_quadro_vira_material():
    pytest.importorskip("rapidocr")

    extraido = asyncio.run(ms.extract(_png_bytes(_LINHAS_AULA), "quadro.jpg"))

    assert extraido.source_type == "image-ocr"
    assert "Normalizacao" in extraido.text
    # Foto de quadro e lista de topicos: as linhas ficam separadas, e nao
    # achatadas num paragrafo corrido.
    assert "\n" in extraido.text


# --- Recorte para o quiz ----------------------------------------------------


def test_recorte_para_o_quiz_termina_em_paragrafo():
    texto = "\n\n".join(_texto_longo(f"Bloco {i}. ") for i in range(1, 8))

    recorte = ms.summary_for_quiz(texto, limit=1200)

    assert len(recorte) <= 1200
    # Corte no meio de uma frase entregaria contexto truncado ao modelo.
    assert not recorte.endswith(("condicao", "a", "de"))


def test_material_pequeno_vai_inteiro():
    texto = "Material curto sobre juncoes."

    assert ms.summary_for_quiz(texto) == texto


# --- Vinculo com a disciplina ------------------------------------------------


class _FakeDisciplineDb:
    """Banco minimo com as disciplinas do professor."""

    def __init__(self, disciplinas):
        self._disciplinas = disciplinas

    async def get(self, _model, item_id):
        return next((d for d in self._disciplinas if d.id == item_id), None)

    async def execute(self, _query):
        disciplinas = self._disciplinas

        class _R:
            def scalars(self):
                class _S:
                    def all(self_inner):
                        return disciplinas

                return _S()

        return _R()


def _disciplina(id_, code, name, tutor="tutor-1"):
    from types import SimpleNamespace

    return SimpleNamespace(id=id_, code=code, name=name, tutor_id=tutor)


def test_material_por_texto_encontra_a_disciplina_cadastrada():
    """Texto que casa com o cadastro ganha o vinculo: sem o id, renomear a
    disciplina soltaria o material dela."""
    from app.routers.education import _resolve_discipline

    db = _FakeDisciplineDb([_disciplina("d1", "ARA0040", "BANCO DE DADOS")])

    vinculo, rotulo = asyncio.run(
        _resolve_discipline("", "ARA0040 - BANCO DE DADOS", "tutor-1", db)
    )

    assert vinculo == "d1"
    assert rotulo == "ARA0040 - BANCO DE DADOS"


def test_material_de_disciplina_nao_cadastrada_ainda_sobe():
    # Travar o upload obrigaria a cadastrar a disciplina antes; o material vale
    # por si, e o vinculo pode vir depois.
    from app.routers.education import _resolve_discipline

    db = _FakeDisciplineDb([])

    vinculo, rotulo = asyncio.run(
        _resolve_discipline("", "Materia nova", "tutor-1", db)
    )

    assert vinculo is None
    assert rotulo == "Materia nova"


def test_disciplina_de_outro_professor_e_recusada():
    from fastapi import HTTPException

    from app.routers.education import _resolve_discipline

    db = _FakeDisciplineDb([_disciplina("d1", "ARA0040", "BD", tutor="outro")])

    with pytest.raises(HTTPException) as erro:
        asyncio.run(_resolve_discipline("d1", "", "tutor-1", db))

    assert erro.value.status_code == 404
