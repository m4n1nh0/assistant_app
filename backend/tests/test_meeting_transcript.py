"""Importacao da transcricao pronta de uma reuniao online.

O Teams entrega legenda (.vtt) com um bloco por frase e o nome em marcacao; o
Meet entrega um documento com "Nome: fala" e horarios soltos. Os dois precisam
chegar ao mesmo lugar: falas com o nome de quem falou, em blocos do tamanho de
um trecho gravado.
"""

from __future__ import annotations

import asyncio
import tempfile

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.database import (
    ClassGroupModel,
    DisciplineModel,
    LessonClassGroupModel,
    LessonModel,
    LessonPointModel,
    LessonSegmentModel,
    ProjectGroupMemberModel,
    ProjectGroupModel,
    StudentModel,
    get_db,
)
from app.core.security import get_current_user
from app.routers import education
from app.services import meeting_transcript_service as service
from app.services.user_llm_config_service import user_llm_context

TEAMS_VTT = """WEBVTT

NOTE duracao 00:01:10

a1b2c3d4-0001/12-0
00:00:03.120 --> 00:00:06.480
<v Ana Souza>Bom dia a todos, vamos comecar.</v>

a1b2c3d4-0001/13-0
00:00:06.900 --> 00:00:09.000
<v Ana Souza>A pauta de hoje &eacute; o calendario de provas.</v>

a1b2c3d4-0001/14-0
00:00:10.000 --> 00:00:14.000
<v Bruno Lima>Eu sugiro adiar a segunda avaliacao.</v>
"""

MEET_DOC = """Reuniao de colegiado - Transcricao

00:00:00

Ana Souza: Bom dia a todos, vamos comecar.
A pauta de hoje e o calendario de provas.
Bruno Lima: Eu sugiro adiar a segunda avaliacao.

00:05:00

Ana Souza: Combinado, fica para a semana seguinte.
"""


def test_vtt_do_teams_perde_a_marcacao_e_junta_falas_seguidas():
    turns = service.parse_transcript(TEAMS_VTT)

    assert [turn.speaker for turn in turns] == ["Ana Souza", "Bruno Lima"]
    assert turns[0].text == (
        "Bom dia a todos, vamos comecar. A pauta de hoje é o calendario de provas."
    )
    assert "-->" not in turns[0].text and "a1b2c3d4" not in turns[0].text


def test_documento_do_meet_ignora_horario_e_continua_a_fala_na_linha_seguinte():
    turns = service.parse_transcript(MEET_DOC)

    assert service.speakers_of(turns) == ["Ana Souza", "Bruno Lima"]
    # O titulo do documento vem antes de qualquer nome: fica sem falante.
    assert turns[0] == service.Turn("", "Reuniao de colegiado - Transcricao")
    assert turns[1].text.endswith("o calendario de provas.")
    assert all("00:05:00" not in turn.text for turn in turns)


# O .docx do Teams: nome e horario numa linha, a fala nas seguintes, e os avisos
# de inicio e fim da transcricao no meio do texto.
TEAMS_DOCX_PARAGRAPHS = [
    "BANCO DE DADOS 3001-20260930_190034-Meeting Recording",
    "September 30, 2026, 10:00PM",
    "2h 1m 28s",
    "\nANA SOUZA started transcription",
    "\nANA SOUZA   0:03\nBoa noite, meus caros alunos.\nVamos comecar pelas 8:10",
    "\nBRUNO LIMA   13:25\nTem que fechar o parenteses ali.",
    "\nANA SOUZA   1:02:09\nIsso mesmo.",
    "\nANA SOUZA stopped transcription",
]


def test_docx_do_teams_reconhece_nome_seguido_de_horario():
    # Como o texto chega colado: tres espacos entre nome e horario.
    turns = service.parse_transcript("\n".join(TEAMS_DOCX_PARAGRAPHS))

    assert turns[1:] == [
        service.Turn("ANA SOUZA", "Boa noite, meus caros alunos. Vamos comecar pelas 8:10"),
        service.Turn("BRUNO LIMA", "Tem que fechar o parenteses ali."),
        service.Turn("ANA SOUZA", "Isso mesmo."),
    ]
    assert "transcription" not in " ".join(turn.text for turn in turns)


def test_docx_do_teams_com_espacos_juntados_ainda_reconhece_o_nome():
    # Como o texto chega do extrator de .docx: paragrafos separados por linha
    # em branco e um espaco so entre nome e horario.
    texto = "\n\n".join(
        " ".join(parte.split(" ")).replace("   ", " ").strip()
        for parte in TEAMS_DOCX_PARAGRAPHS
    )

    turns = service.parse_transcript(texto)

    assert service.speakers_of(turns) == ["ANA SOUZA", "BRUNO LIMA"]
    # Frase que termina em horario, no meio da fala, nao vira falante.
    assert turns[1].text.endswith("Vamos comecar pelas 8:10")


def test_srt_sem_marcacao_de_voz_reconhece_nome_no_texto():
    srt = "1\n00:00:01,000 --> 00:00:03,000\nAna Souza: Bom dia.\n\n2\n" \
          "00:00:03,500 --> 00:00:05,000\nSem nome aqui.\n"

    turns = service.parse_transcript(srt)

    assert turns == [
        service.Turn("Ana Souza", "Bom dia."),
        service.Turn("", "Sem nome aqui."),
    ]


def test_texto_sem_nome_nenhum_vira_uma_fala_so():
    turns = service.parse_transcript("primeira linha\nsegunda linha\n")

    assert turns == [service.Turn("", "primeira linha segunda linha")]


def test_blocos_respeitam_o_tamanho_e_repetem_o_nome_em_fala_longa():
    longa = " ".join(f"Frase numero {i} da explicacao." for i in range(120))
    turns = [service.Turn("Ana Souza", longa), service.Turn("Bruno Lima", "Certo.")]

    blocks = service.split_blocks(turns, max_chars=600)

    assert len(blocks) > 3
    assert all(len(block) <= 600 for block in blocks)
    assert all(block.startswith("Ana Souza: ") for block in blocks[:-1])
    assert blocks[-1].endswith("Bruno Lima: Certo.")
    # Nada da fala se perde no corte.
    recomposto = " ".join(
        linha.split(": ", 1)[1] for bloco in blocks for linha in bloco.split("\n")
    )
    assert recomposto == f"{longa} Certo."


def test_fala_sem_pontuacao_e_cortada_em_espaco():
    turns = [service.Turn("", "palavra " * 400)]

    blocks = service.split_blocks(turns, max_chars=500)

    assert all(len(block) <= 500 for block in blocks)
    assert "".join(blocks).count("palavra") == 400


def test_arquivo_salvo_em_latin1_nao_e_recusado():
    assert service.decode_text("reunião".encode("latin-1")) == "reunião"
    assert service.decode_text("﻿reunião".encode("utf-8")) == "reunião"


# --- Rota --------------------------------------------------------------------

USER = {"uid": "u1", "tutor_id": "t1"}


@pytest.fixture
def client(monkeypatch):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tempfile.mkdtemp()}/aulas.db")
    sessions = async_sessionmaker(engine, expire_on_commit=False)

    async def seed():
        async with engine.begin() as conn:
            for model in (
                LessonModel, LessonClassGroupModel, LessonSegmentModel,
                LessonPointModel, ClassGroupModel, DisciplineModel,
                ProjectGroupModel, ProjectGroupMemberModel, StudentModel,
            ):
                await conn.run_sync(model.__table__.create)

    asyncio.run(seed())

    async def db_dependency():
        async with sessions() as session:
            yield session

    indexados = []

    async def _index(**kwargs):
        indexados.extend(kwargs["segments"])
        return len(kwargs["segments"])

    monkeypatch.setattr(education.qdrant_service, "index_lesson_segments", _index)

    app = FastAPI()
    app.include_router(education.router)
    app.dependency_overrides[get_db] = db_dependency
    app.dependency_overrides[get_current_user] = lambda: USER
    app.dependency_overrides[user_llm_context] = lambda: None
    with TestClient(app) as test_client:
        test_client.indexados = indexados
        yield test_client
    asyncio.run(engine.dispose())


def _palestra(client, title="Reuniao de colegiado"):
    return client.post(
        "/education/lessons", json={"kind": "palestra", "title": title}
    ).json()["id"]


pytestmark_rota = pytest.mark.integration


@pytestmark_rota
def test_texto_colado_vira_trechos_da_gravacao(client):
    lesson_id = _palestra(client)

    resposta = client.post(
        f"/education/lessons/{lesson_id}/transcript", data={"text": MEET_DOC}
    )

    assert resposta.status_code == 200, resposta.text
    corpo = resposta.json()
    assert corpo["imported"] == 1 and corpo["indexed"] == 1
    assert corpo["speakers"] == ["Ana Souza", "Bruno Lima"]
    assert corpo["lesson"]["segment_count"] == 1

    detalhe = client.get(f"/education/lessons/{lesson_id}").json()
    texto = detalhe["segments"][0]["text"]
    assert "Bruno Lima: Eu sugiro adiar a segunda avaliacao." in texto
    assert "00:05:00" not in texto
    assert client.indexados[0]["text"] == texto


@pytestmark_rota
def test_arquivo_vtt_do_teams_e_aceito(client):
    lesson_id = _palestra(client)

    resposta = client.post(
        f"/education/lessons/{lesson_id}/transcript",
        files={"file": ("reuniao.vtt", TEAMS_VTT.encode("utf-8"), "text/vtt")},
    )

    assert resposta.status_code == 200, resposta.text
    assert resposta.json()["speakers"] == ["Ana Souza", "Bruno Lima"]


@pytestmark_rota
def test_arquivo_docx_do_teams_e_aceito(client):
    import io

    from docx import Document

    documento = Document()
    for paragrafo in TEAMS_DOCX_PARAGRAPHS:
        documento.add_paragraph(paragrafo)
    # O extrator recusa documento com menos de 200 caracteres.
    documento.add_paragraph("\nBRUNO LIMA   1:03:00\n" + "Mais uma fala longa. " * 12)
    arquivo = io.BytesIO()
    documento.save(arquivo)
    lesson_id = _palestra(client)

    resposta = client.post(
        f"/education/lessons/{lesson_id}/transcript",
        files={"file": ("reuniao.docx", arquivo.getvalue(), "application/octet-stream")},
    )

    assert resposta.status_code == 200, resposta.text
    assert resposta.json()["speakers"] == ["ANA SOUZA", "BRUNO LIMA"]
    texto = client.get(f"/education/lessons/{lesson_id}").json()["segments"][0]["text"]
    assert "ANA SOUZA: Boa noite, meus caros alunos." in texto
    assert "0:03" not in texto and "transcription" not in texto


@pytestmark_rota
def test_transcricao_longa_e_dividida_em_varios_trechos(client):
    lesson_id = _palestra(client)
    texto = "\n".join(
        f"{'Ana Souza' if i % 2 else 'Bruno Lima'}: fala numero {i} " + "bla " * 60
        for i in range(30)
    )

    corpo = client.post(
        f"/education/lessons/{lesson_id}/transcript", data={"text": texto}
    ).json()

    assert corpo["imported"] > 3
    assert corpo["lesson"]["segment_count"] == corpo["imported"]
    sequencias = [item["sequence"] for item in client.indexados]
    assert sequencias == list(range(1, corpo["imported"] + 1))


@pytestmark_rota
def test_formato_desconhecido_texto_vazio_e_aula_encerrada_sao_recusados(client):
    lesson_id = _palestra(client)
    rota = f"/education/lessons/{lesson_id}/transcript"

    assert client.post(
        rota, files={"file": ("reuniao.mp4", b"\x00\x01", "video/mp4")}
    ).status_code == 422
    assert client.post(rota, data={"text": "   \n00:00:00\n"}).status_code == 422

    client.post(f"/education/lessons/{lesson_id}/status", json={"status": "closed"})
    assert client.post(rota, data={"text": MEET_DOC}).status_code == 409


@pytestmark_rota
def test_gravacao_de_outro_professor_nao_recebe_transcricao(client):
    assert client.post(
        "/education/lessons/nao-e-minha/transcript", data={"text": MEET_DOC}
    ).status_code == 404
