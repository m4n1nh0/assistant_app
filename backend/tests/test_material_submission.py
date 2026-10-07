"""Link publico para os alunos enviarem o material da apresentacao do grupo.

Falha que isto evita: o material das apresentacoes chegava por e-mail e chat, ficava
solto do grupo e nao dava para usar no quiz. Aqui o aluno entra com a matricula, o
arquivo vira material ligado ao grupo dele (inclusive nos grupos da aula reunida) e o
professor ve, por grupo, o que chegou e o que foi gravado.
"""

from __future__ import annotations

import asyncio
import io
import tempfile
from datetime import datetime, timedelta, timezone

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.database import (
    ClassGroupModel,
    ClassScheduleModel,
    DisciplineModel,
    LessonModel,
    LessonSegmentModel,
    MaterialModel,
    MaterialSubmissionLinkModel,
    ProjectGroupClassModel,
    ProjectGroupPointModel,
    ProjectGroupMemberModel,
    ProjectGroupModel,
    StudentModel,
    get_db,
)
from app.core.security import get_current_user
from app.routers import education, material_submission
from app.services import material_submission_service as submissions

USER = {"uid": "u1", "tutor_id": "t1"}
TABELAS = (
    DisciplineModel, ClassGroupModel, ClassScheduleModel, StudentModel,
    ProjectGroupModel, ProjectGroupClassModel, ProjectGroupPointModel,
    ProjectGroupMemberModel,
    MaterialModel, MaterialSubmissionLinkModel, LessonModel, LessonSegmentModel,
)


def pptx_bytes(slides=(("Banco de dados do projeto", [
        "Modelo relacional com tabelas de alunos, grupos e apresentacoes do semestre",
        "Normalizacao ate a 3FN para evitar redundancia entre as tabelas do sistema",
        "Consultas com juncao interna para listar os grupos com seus integrantes",
]),)):
    from pptx import Presentation

    deck = Presentation()
    layout = deck.slide_layouts[1]
    for titulo, bullets in slides:
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = titulo
        corpo = slide.placeholders[1].text_frame
        corpo.text = bullets[0]
        for bullet in bullets[1:]:
            corpo.add_paragraph().text = bullet
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def pptx_sem_texto() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def aluno(id_, nome, turma, matricula):
    return StudentModel(id=id_, tutor_id="t1", name=nome, class_id=turma,
                        class_group="x", external_id=matricula, active=True)


@pytest.fixture
def api():
    engine = create_async_engine(f"sqlite+aiosqlite:///{tempfile.mkdtemp()}/m.db")
    sessions = async_sessionmaker(engine, expire_on_commit=False)

    async def seed():
        async with engine.begin() as conn:
            for model in TABELAS:
                await conn.run_sync(model.__table__.create)
        async with sessions() as db:
            db.add_all([
                DisciplineModel(id="d1", tutor_id="t1", code="ARA0040", name="BANCO DE DADOS", semester="2026.2"),
                DisciplineModel(id="d2", tutor_id="t1", code="ARA0058", name="CLOUD", semester="2026.2"),
                ClassGroupModel(id="c-3002", tutor_id="t1", code="3002", name="A", discipline_id="d1", discipline="BD"),
                ClassGroupModel(id="c-3030", tutor_id="t1", code="3030", name="B", discipline_id="d1", discipline="BD"),
                ClassGroupModel(id="c-qui", tutor_id="t1", code="3001", name="Q", discipline_id="d1", discipline="BD"),
                ClassGroupModel(id="c-cloud", tutor_id="t1", code="4001", name="C", discipline_id="d2", discipline="CLOUD"),
                aluno("s-ana", "Ana Souza", "c-3002", "2024-0001"),
                aluno("s-bia", "Bia Lima", "c-3030", "20240002"),
                aluno("s-caio", "Caio Reis", "c-3030", "20240003"),
                aluno("s-davi", "Davi Alves", "c-qui", "20240004"),
                aluno("s-eva", "Eva Costa", "c-3002", "20240005"),  # sem grupo
                ProjectGroupModel(id="g1", tutor_id="t1", discipline_id="d1", semester="2026.2", name="GRUPO 1", class_id="c-3002"),
                ProjectGroupModel(id="g2", tutor_id="t1", discipline_id="d1", semester="2026.2", name="GRUPO 2", class_id="c-3030"),
                ProjectGroupModel(id="g3", tutor_id="t1", discipline_id="d1", semester="2026.2", name="GRUPO 3", class_id="c-qui"),
                ProjectGroupClassModel(group_id="g1", class_id="c-3002"),
                ProjectGroupClassModel(group_id="g1", class_id="c-3030"),
                ProjectGroupMemberModel(id="m-ana", group_id="g1", student_id="s-ana", name="Ana Souza", position=0),
                ProjectGroupMemberModel(id="m-bia", group_id="g1", student_id="s-bia", name="Bia Lima", position=1),
                ProjectGroupMemberModel(id="m-caio", group_id="g2", student_id="s-caio", name="Caio Reis", position=0),
                ProjectGroupMemberModel(id="m-davi", group_id="g3", student_id="s-davi", name="Davi Alves", position=0),
            ])
            await db.commit()

    asyncio.run(seed())

    async def db_dependency():
        async with sessions() as session:
            yield session

    app = FastAPI()
    app.include_router(education.router)
    app.include_router(material_submission.router)
    app.dependency_overrides[get_db] = db_dependency
    app.dependency_overrides[get_current_user] = lambda: USER
    with TestClient(app) as client:
        client.sessions = sessions
        yield client
    asyncio.run(engine.dispose())


def criar_link(api, **extra):
    corpo = {"discipline_id": "d1", "title": "Slides do projeto", **extra}
    resposta = api.post("/education/material-links", json=corpo)
    assert resposta.status_code == 200, resposta.text
    return resposta.json()


def enviar(api, link, matricula, nome="slides.pptx", dados=None, titulo=""):
    return api.post(
        f"/education/material-submit/{link['token']}/send",
        data={"enrollment": matricula, "title": titulo},
        files={"file": (nome, dados if dados is not None else pptx_bytes(),
                        "application/octet-stream")},
    )


def buscar(api, link, matricula):
    return api.post(f"/education/material-submit/{link['token']}/lookup",
                    json={"enrollment": matricula})


def materiais(api, **params):
    return api.get("/education/materials", params=params).json()


# --- link (professor) ------------------------------------------------------------


def test_cria_o_link_com_token_caminho_e_contagem_de_grupos(api):
    link = criar_link(api)

    assert len(link["token"]) >= 12
    assert link["path"] == f"/education/material-submit/{link['token']}"
    assert link["state"] == "open" and link["active"] is True
    assert link["groups_total"] == 3 and link["groups_sent"] == 0
    assert link["discipline"] == "ARA0040 - BANCO DE DADOS"
    assert link["max_files_per_group"] == 5


def test_link_por_turmas_conta_so_os_grupos_delas(api):
    link = criar_link(api, class_ids=["c-3030", "c-3002"])

    assert link["class_ids"] == ["c-3002", "c-3030"]
    assert link["groups_total"] == 2
    assert " + " in link["class_label"]


def test_turma_de_outra_disciplina_ou_disciplina_alheia_e_recusada(api):
    assert api.post("/education/material-links", json={
        "discipline_id": "d1", "class_ids": ["c-cloud"]}).status_code == 422
    assert api.post("/education/material-links", json={
        "discipline_id": "nao-existe"}).status_code == 404


def test_listar_filtra_por_disciplina(api):
    criar_link(api)

    assert len(api.get("/education/material-links").json()) == 1
    assert api.get("/education/material-links", params={"discipline_id": "d2"}).json() == []


def test_fechar_reabrir_prazo_e_limite(api):
    link = criar_link(api)
    base = f"/education/material-links/{link['id']}"

    fechado = api.patch(base, json={"active": False}).json()
    assert fechado["state"] == "closed"

    reaberto = api.patch(base, json={"active": True, "max_files_per_group": 2}).json()
    assert reaberto["state"] == "open" and reaberto["max_files_per_group"] == 2

    prazo = (datetime.now(timezone.utc) - timedelta(minutes=1)).isoformat()
    vencido = api.patch(base, json={"closes_at": prazo}).json()
    assert vencido["state"] == "expired"

    sem_prazo = api.patch(base, json={"clear_deadline": True}).json()
    assert sem_prazo["state"] == "open" and sem_prazo["closes_at"] is None


def test_apagar_o_link_derruba_a_url_mas_mantem_o_material(api):
    link = criar_link(api)
    assert enviar(api, link, "20240002").status_code == 200

    assert api.delete(f"/education/material-links/{link['id']}").status_code == 200

    assert api.get(f"/education/material-submit/{link['token']}").status_code == 200  # pagina amigavel
    assert buscar(api, link, "20240002").status_code == 404
    assert len(materiais(api)) == 1


# --- pagina publica ---------------------------------------------------------------


def test_pagina_publica_mostra_titulo_disciplina_e_nao_vaza_o_token_sem_escape(api):
    link = criar_link(api, title="Slides <b>finais</b>")

    pagina = api.get(f"/education/material-submit/{link['token']}")

    assert pagina.status_code == 200
    assert "Slides &lt;b&gt;finais&lt;/b&gt;" in pagina.text
    assert "<b>finais</b>" not in pagina.text
    assert "ARA0040 - BANCO DE DADOS" in pagina.text
    assert "BUSCAR MEU GRUPO" in pagina.text
    assert "noindex" in pagina.text


def test_pagina_de_link_fechado_vencido_ou_desconhecido_explica(api):
    link = criar_link(api)
    base = f"/education/material-links/{link['id']}"

    assert "não existe" in api.get("/education/material-submit/nao-existe").text

    api.patch(base, json={"active": False})
    assert "encerrou" in api.get(f"/education/material-submit/{link['token']}").text

    api.patch(base, json={"active": True, "closes_at":
              (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat()})
    assert "prazo" in api.get(f"/education/material-submit/{link['token']}").text


# --- identificar o grupo ------------------------------------------------------------


def test_matricula_acha_o_grupo_e_os_integrantes(api):
    link = criar_link(api)

    resposta = buscar(api, link, "2024-0001")  # pontuacao ignorada, como no quiz

    assert resposta.status_code == 200, resposta.text
    dados = resposta.json()
    assert dados["group"]["name"] == "GRUPO 1"
    assert dados["member"]["name"] == "Ana Souza"
    assert dados["members"] == ["Ana Souza", "Bia Lima"]
    assert dados["materials"] == []


def test_aluno_da_outra_turma_da_aula_reunida_acha_o_mesmo_grupo(api):
    link = criar_link(api, class_ids=["c-3002", "c-3030"])

    ana = buscar(api, link, "20240001").json()
    bia = buscar(api, link, "20240002").json()

    assert ana["group"]["id"] == bia["group"]["id"] == "g1"


def test_matricula_desconhecida_ou_sem_grupo_ou_de_turma_fora_do_link(api):
    link = criar_link(api, class_ids=["c-3002", "c-3030"])

    assert buscar(api, link, "").json()["code"] == "invalid"
    desconhecida = buscar(api, link, "99999999")
    assert desconhecida.status_code == 404 and desconhecida.json()["code"] == "unknown"
    # Eva existe na disciplina mas nao esta em grupo; Davi e de grupo de outra turma.
    assert buscar(api, link, "20240005").json()["code"] == "no_group"
    assert buscar(api, link, "20240004").json()["code"] == "no_group"


def test_lookup_em_link_fechado_responde_409(api):
    link = criar_link(api)
    api.patch(f"/education/material-links/{link['id']}", json={"active": False})

    resposta = buscar(api, link, "20240001")

    assert resposta.status_code == 409 and resposta.json()["code"] == "closed"


# --- enviar ----------------------------------------------------------------------


def test_envio_vira_material_do_grupo_com_o_texto_extraido(api):
    link = criar_link(api)

    resposta = enviar(api, link, "20240002", titulo="Slides finais")

    assert resposta.status_code == 200, resposta.text
    dados = resposta.json()
    assert dados["group"] == "GRUPO 1" and dados["replaced"] is False
    assert dados["material"]["title"] == "Slides finais"

    lista = materiais(api, group_id="g1")
    assert len(lista) == 1
    item = lista[0]
    assert item["group_id"] == "g1" and item["from_link"] is True
    assert item["uploader_name"] == "Bia Lima"
    assert item["source_type"] == "pptx"
    assert item["discipline_id"] == "d1"

    async def texto():
        async with api.sessions() as db:
            return (await db.get(MaterialModel, item["id"])).content
    conteudo = asyncio.run(texto())
    assert "Normalizacao ate a 3FN" in conteudo


def test_sem_titulo_usa_o_do_documento_ou_o_nome_do_arquivo(api):
    link = criar_link(api)

    enviar(api, link, "20240003", nome="entrega final.pptx")

    assert materiais(api, group_id="g2")[0]["title"]


def test_reenviar_o_mesmo_arquivo_substitui_em_vez_de_duplicar(api):
    link = criar_link(api)
    corpo = " com bastante texto para passar do minimo exigido pelo leitor de material " * 4
    enviar(api, link, "20240001",
           dados=pptx_bytes([("Versao 1", ["primeira versao do texto" + corpo])]))

    segunda = enviar(api, link, "20240002",
                     dados=pptx_bytes([("Versao 2", ["segunda versao do texto" + corpo])]))

    assert segunda.json()["replaced"] is True
    lista = materiais(api, group_id="g1")
    assert len(lista) == 1

    async def texto():
        async with api.sessions() as db:
            return (await db.get(MaterialModel, lista[0]["id"])).content
    assert "segunda versao" in asyncio.run(texto())
    assert "primeira versao" not in asyncio.run(texto())


def test_limite_de_arquivos_por_grupo(api):
    link = criar_link(api, max_files_per_group=2)
    assert enviar(api, link, "20240001", nome="a.pptx").status_code == 200
    assert enviar(api, link, "20240001", nome="b.pptx").status_code == 200

    terceiro = enviar(api, link, "20240001", nome="c.pptx")

    assert terceiro.status_code == 409 and terceiro.json()["code"] == "limit"
    # Trocar um que ja existe continua valendo.
    assert enviar(api, link, "20240001", nome="b.pptx").status_code == 200
    # O outro grupo tem o seu proprio limite.
    assert enviar(api, link, "20240003", nome="a.pptx").status_code == 200


@pytest.mark.parametrize("nome, codigo", [
    ("slides.ppt", "format"),
    ("notas.txt", "format"),
    ("foto.png", "format"),
    ("sem_extensao", "format"),
])
def test_formato_nao_aceito(api, nome, codigo):
    link = criar_link(api)

    resposta = enviar(api, link, "20240001", nome=nome, dados=b"qualquer coisa")

    assert resposta.status_code == 422 and resposta.json()["code"] == codigo
    assert materiais(api) == []


def test_ppt_antigo_explica_como_converter(api):
    link = criar_link(api)

    detalhe = enviar(api, link, "20240001", nome="x.ppt", dados=b"x").json()["detail"]

    assert "pptx" in detalhe.lower()


def test_arquivo_vazio_corrompido_ou_sem_texto(api):
    link = criar_link(api)

    assert enviar(api, link, "20240001", dados=b"").json()["code"] == "file"
    corrompido = enviar(api, link, "20240001", dados=b"isto nao e um pptx")
    assert corrompido.status_code == 422 and corrompido.json()["code"] == "extract"
    sem_texto = enviar(api, link, "20240001", dados=pptx_sem_texto())
    assert sem_texto.status_code == 422 and sem_texto.json()["code"] == "empty"
    assert materiais(api) == []


def test_arquivo_grande_demais(api, monkeypatch):
    monkeypatch.setattr(submissions, "MAX_UPLOAD_BYTES", 1024)
    link = criar_link(api)

    resposta = enviar(api, link, "20240001", dados=pptx_bytes())

    assert resposta.status_code == 413 and resposta.json()["code"] == "size"


def test_envio_com_matricula_errada_nao_grava(api):
    link = criar_link(api)

    resposta = enviar(api, link, "00000000")

    assert resposta.status_code == 404 and resposta.json()["code"] == "unknown"
    assert materiais(api) == []


def test_envio_em_link_fechado_ou_vencido_e_recusado(api):
    link = criar_link(api)
    base = f"/education/material-links/{link['id']}"

    api.patch(base, json={"active": False})
    assert enviar(api, link, "20240001").json()["code"] == "closed"

    api.patch(base, json={"active": True, "closes_at":
              (datetime.now(timezone.utc) - timedelta(minutes=5)).isoformat()})
    assert enviar(api, link, "20240001").json()["code"] == "expired"
    assert materiais(api) == []


def test_link_com_turmas_nao_aceita_grupo_de_fora(api):
    link = criar_link(api, class_ids=["c-3002", "c-3030"])

    assert enviar(api, link, "20240004").json()["code"] == "no_group"  # Davi, da quinta


# --- remover -----------------------------------------------------------------------


def test_aluno_remove_so_arquivo_do_proprio_grupo_pelo_mesmo_link(api):
    link = criar_link(api)
    enviar(api, link, "20240001", nome="g1.pptx")
    enviar(api, link, "20240003", nome="g2.pptx")
    g1 = materiais(api, group_id="g1")[0]["id"]
    g2 = materiais(api, group_id="g2")[0]["id"]
    url = f"/education/material-submit/{link['token']}/remove"

    alheio = api.post(url, json={"enrollment": "20240001", "material_id": g2})
    assert alheio.status_code == 404 and alheio.json()["code"] == "not_found"

    proprio = api.post(url, json={"enrollment": "20240002", "material_id": g1})
    assert proprio.status_code == 200
    assert materiais(api, group_id="g1") == []
    assert len(materiais(api, group_id="g2")) == 1


def test_nao_remove_material_do_professor_nem_de_outro_link(api):
    link = criar_link(api)
    outro = criar_link(api, title="Outro link")
    enviar(api, outro, "20240001", nome="do-outro.pptx")
    item = materiais(api, group_id="g1")[0]["id"]

    resposta = api.post(f"/education/material-submit/{link['token']}/remove",
                        json={"enrollment": "20240001", "material_id": item})

    assert resposta.status_code == 404
    assert len(materiais(api, group_id="g1")) == 1


# --- visao do professor ---------------------------------------------------------------


async def _semear_gravacao(sessions, grupo="g1", segmentos=2, kind="apresentacao"):
    async with sessions() as db:
        aula = LessonModel(id="l1", tutor_id="t1", kind=kind, group_id=grupo,
                           discipline="ARA0040 - BANCO DE DADOS", semester="2026.2",
                           title="Apresentacao: GRUPO 1", status="closed")
        db.add(aula)
        for numero in range(segmentos):
            db.add(LessonSegmentModel(lesson_id="l1", tutor_id="t1", sequence=numero,
                                      text=f"trecho {numero}"))
        await db.commit()


def test_painel_mostra_material_gravacao_e_o_que_falta(api):
    link = criar_link(api)
    enviar(api, link, "20240001")
    asyncio.run(_semear_gravacao(api.sessions))

    resposta = api.get("/education/presentations", params={"discipline_id": "d1"})

    assert resposta.status_code == 200, resposta.text
    por_nome = {item["group_name"]: item for item in resposta.json()}
    g1 = por_nome["GRUPO 1"]
    assert len(g1["materials"]) == 1 and g1["materials"][0]["uploader_name"] == "Ana Souza"
    assert g1["materials"][0]["from_link"] is True
    assert len(g1["recordings"]) == 1 and g1["recordings"][0]["segments"] == 2
    assert g1["gaps"] == []
    assert g1["class_ids"] == ["c-3002", "c-3030"]
    assert por_nome["GRUPO 2"]["gaps"] == ["material", "gravação"]


def test_gravacao_sem_transcricao_ainda_conta_como_falta(api):
    link = criar_link(api)
    enviar(api, link, "20240001")
    asyncio.run(_semear_gravacao(api.sessions, segmentos=0))

    g1 = next(item for item in api.get(
        "/education/presentations", params={"discipline_id": "d1"}).json()
        if item["group_name"] == "GRUPO 1")

    assert g1["gaps"] == ["gravação"]


def test_painel_filtra_por_turma_e_valida_a_disciplina(api):
    todos = api.get("/education/presentations", params={"discipline_id": "d1"}).json()
    quinta = api.get("/education/presentations",
                     params={"discipline_id": "d1", "class_id": "c-qui"}).json()

    assert len(todos) == 3
    assert [item["group_name"] for item in quinta] == ["GRUPO 3"]
    assert api.get("/education/presentations",
                   params={"discipline_id": "nao-existe"}).status_code == 404


def test_progresso_do_link_conta_arquivos_e_grupos(api):
    link = criar_link(api)
    enviar(api, link, "20240001", nome="a.pptx")
    enviar(api, link, "20240002", nome="b.pptx")
    enviar(api, link, "20240003", nome="c.pptx")

    atual = api.get("/education/material-links").json()[0]

    assert atual["files_received"] == 3
    assert atual["groups_sent"] == 2 and atual["groups_total"] == 3


# --- ligar material e gravacao ao grupo -------------------------------------------


def test_liga_um_material_do_professor_a_um_grupo_e_solta(api):
    async def semear():
        async with api.sessions() as db:
            db.add(MaterialModel(id="mat1", tutor_id="t1", discipline_id="d1",
                                 discipline="ARA0040 - BANCO DE DADOS", title="Slides do G3",
                                 filename="g3.pdf", source_type="pdf", content="texto"))
            await db.commit()
    asyncio.run(semear())

    ligado = api.put("/education/materials/mat1/group", json={"group_id": "g3"})
    assert ligado.status_code == 200, ligado.text
    assert ligado.json()["group_id"] == "g3"
    assert len(materiais(api, group_id="g3")) == 1

    solto = api.put("/education/materials/mat1/group", json={"group_id": None})
    assert solto.json()["group_id"] is None


def test_material_nao_liga_a_grupo_de_outra_disciplina_ou_alheio(api):
    async def semear():
        async with api.sessions() as db:
            db.add(MaterialModel(id="mat1", tutor_id="t1", discipline_id="d2",
                                 discipline="CLOUD", title="X", filename="x.pdf",
                                 source_type="pdf", content="texto"))
            await db.commit()
    asyncio.run(semear())

    assert api.put("/education/materials/mat1/group",
                   json={"group_id": "g1"}).status_code == 422
    assert api.put("/education/materials/mat1/group",
                   json={"group_id": "nao-existe"}).status_code == 404
    assert api.put("/education/materials/nao-existe/group",
                   json={"group_id": "g1"}).status_code == 404


def test_liga_uma_palestra_ao_grupo_e_ela_vira_apresentacao(api):
    asyncio.run(_semear_gravacao(api.sessions, grupo=None, kind="palestra"))

    resposta = api.put("/education/lessons/l1/presentation-group", json={"group_id": "g2"})

    assert resposta.status_code == 200, resposta.text
    dados = resposta.json()
    assert dados["kind"] == "apresentacao" and dados["group_id"] == "g2"
    assert dados["group_name"] == "GRUPO 2"
    g2 = next(item for item in api.get(
        "/education/presentations", params={"discipline_id": "d1"}).json()
        if item["group_name"] == "GRUPO 2")
    assert len(g2["recordings"]) == 1


def test_soltar_a_gravacao_volta_a_palestra(api):
    asyncio.run(_semear_gravacao(api.sessions))

    resposta = api.put("/education/lessons/l1/presentation-group", json={"group_id": None})

    assert resposta.json()["kind"] == "palestra" and resposta.json()["group_id"] is None


def test_aula_nao_se_liga_a_grupo(api):
    asyncio.run(_semear_gravacao(api.sessions, grupo=None, kind="aula"))

    resposta = api.put("/education/lessons/l1/presentation-group", json={"group_id": "g1"})

    assert resposta.status_code == 422


def test_ligar_gravacao_a_grupo_alheio_ou_inexistente(api):
    asyncio.run(_semear_gravacao(api.sessions, grupo=None, kind="palestra"))

    assert api.put("/education/lessons/l1/presentation-group",
                   json={"group_id": "nao-existe"}).status_code == 404
    assert api.put("/education/lessons/nao-existe/presentation-group",
                   json={"group_id": "g1"}).status_code == 404
