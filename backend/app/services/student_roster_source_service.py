"""Prévia de listas de alunos vindas de planilha, página copiada ou print."""

from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
from pathlib import Path


MAX_SOURCE_BYTES = 10_000_000
MAX_SOURCE_ROWS = 3000
ENROLLMENT_HEADERS = {
    "matricula", "matriculaaluno", "numeromatricula", "ra", "registro",
    "registroacademico", "codigoaluno", "idaluno", "idmatricula",
}
NAME_HEADERS = {
    "nome", "nomecompleto", "nomealuno", "nomedoaluno", "aluno",
    "estudante", "discente", "nomeestudante",
}
CLASS_HEADERS = {"turma", "codigoturma", "codigodaturma", "sequenciaturma",
                 "numeroturma", "classgroup"}
DISCIPLINE_HEADERS = {"disciplina", "codigodisciplina", "materia",
                      "componentecurricular"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}


def _clean_header(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value).casefold())
    return re.sub(r"[^a-z0-9]", "", "".join(character for character in value
        if not unicodedata.combining(character)))


def _looks_enrollment(value: str) -> bool:
    return bool(re.fullmatch(r"\d{7,16}", value.strip()))


def _looks_name(value: str) -> bool:
    return bool(re.fullmatch(r"[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ .'-]{3,178}", value.strip())
                and len(value.strip().split()) >= 2)


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _delimited_rows(text: str) -> list[list[str]]:
    sample = "\n".join(text.splitlines()[:5])
    delimiter = max(("\t", ";", ",", "|"), key=lambda item: sample.count(item))
    if sample.count(delimiter) == 0:
        return []
    return [[cell.strip() for cell in row] for row in
            csv.reader(io.StringIO(text), delimiter=delimiter)
            if any(cell.strip() for cell in row)][:MAX_SOURCE_ROWS + 1]


def _line_pair_rows(text: str) -> list[list[str]]:
    rows = []
    pending_enrollment = ""
    pending_name = ""
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        enrollment = re.search(r"(?<!\d)\d{7,16}(?!\d)", line)
        if enrollment:
            name = (line[:enrollment.start()] + " " + line[enrollment.end():]).strip()
            name = re.sub(r"\b(matr[ií]cula|ra|registro)\s*:?", "", name,
                          flags=re.IGNORECASE).strip(" :-")
            if _looks_name(name):
                rows.append([enrollment.group(), name])
                pending_enrollment = pending_name = ""
            elif _looks_name(pending_name):
                rows.append([enrollment.group(), pending_name])
                pending_enrollment = pending_name = ""
            else:
                pending_enrollment = enrollment.group()
        elif _looks_name(line) and pending_enrollment:
            rows.append([pending_enrollment, line])
            pending_enrollment = pending_name = ""
        elif _looks_name(line):
            pending_name = line
        if len(rows) >= MAX_SOURCE_ROWS:
            break
    return rows


class _TableHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows: list[list[str]] = []
        self.row: list[str] | None = None
        self.cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in {"td", "th"} and self.cell is not None and self.row is not None:
            self.row.append(" ".join(" ".join(self.cell).split()))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            if any(self.row):
                self.rows.append(self.row)
            self.row = None


def _xml_rows(text: str) -> list[list[str]]:
    if re.search(r"<!\s*(DOCTYPE|ENTITY)", text, re.IGNORECASE):
        raise ValueError("XML com DTD ou entidades não é aceito")
    root = ET.fromstring(text)
    records = []
    for element in root.iter():
        fields = {**{_clean_header(key.split("}")[-1]): value
                     for key, value in element.attrib.items()}}
        for child in element:
            if len(child) == 0 and child.text and child.text.strip():
                fields[_clean_header(child.tag.split("}")[-1])] = child.text.strip()
        if any(key in ENROLLMENT_HEADERS for key in fields) and any(
                key in NAME_HEADERS for key in fields):
            records.append(fields)
        if len(records) >= MAX_SOURCE_ROWS:
            break
    if not records:
        raise ValueError("Não encontrei registros com matrícula e nome no XML")
    columns = list(dict.fromkeys(key for row in records for key in row))
    return [columns, *[[row.get(column, "") for column in columns]
                       for row in records]]


def _json_rows(text: str) -> list[list[str]]:
    data = json.loads(text)
    if isinstance(data, dict):
        data = next((value for value in data.values() if isinstance(value, list)), [])
    if not isinstance(data, list) or not all(isinstance(row, dict) for row in data):
        raise ValueError("JSON precisa conter uma lista de alunos")
    records = data[:MAX_SOURCE_ROWS]
    columns = list(dict.fromkeys(key for row in records for key in row))
    return [columns, *[[str(row.get(column, "") or "") for column in columns]
                       for row in records]]


def _xlsx_rows(data: bytes) -> tuple[list[list[str]], str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Leitura de XLSX indisponível neste servidor") from exc
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = [[str(cell).strip() if cell is not None else "" for cell in row]
                    for row in sheet.iter_rows(values_only=True, max_row=MAX_SOURCE_ROWS + 1)]
            rows = [row for row in rows if any(row)]
            if rows:
                return rows, sheet.title
    finally:
        workbook.close()
    raise ValueError("Planilha sem linhas para importar")


def _mapping(rows: list[list[str]]) -> tuple[list[str], list[list[str]], int, int, list[str]]:
    width = min(max(map(len, rows)), 30)
    first = [rows[0][index] if index < len(rows[0]) else "" for index in range(width)]
    headers = [_clean_header(value) for value in first]
    enrollment = next((index for index, value in enumerate(headers)
                       if value in ENROLLMENT_HEADERS), -1)
    name = next((index for index, value in enumerate(headers)
                 if value in NAME_HEADERS), -1)
    has_header = enrollment >= 0 or name >= 0
    columns = first if has_header else [f"Coluna {index + 1}" for index in range(width)]
    data_rows = rows[1:] if has_header else rows
    warnings = []
    if enrollment < 0:
        scores = [sum(_looks_enrollment(row[index]) for row in data_rows[:30]
                      if index < len(row)) for index in range(width)]
        if scores and max(scores) >= max(2, len(data_rows[:30]) // 2):
            enrollment = scores.index(max(scores))
    if name < 0:
        scores = [sum(_looks_name(row[index]) for row in data_rows[:30]
                      if index < len(row)) for index in range(width)]
        ranked = sorted(range(width), key=lambda index: -scores[index])
        name = next((index for index in ranked if index != enrollment and
                     scores[index] >= max(2, len(data_rows[:30]) // 2)), -1)
    if enrollment < 0 or name < 0 or enrollment == name:
        warnings.append("Confira quais colunas contêm matrícula e nome antes de importar")
    return columns, data_rows[:MAX_SOURCE_ROWS], enrollment, name, warnings


def _scope_hints(columns: list[str], rows: list[list[str]],
                 enrollment_index: int) -> dict:
    headers = [_clean_header(column) for column in columns]
    class_index = next((index for index, value in enumerate(headers)
                        if value in CLASS_HEADERS), -1)
    discipline_index = next((index for index, value in enumerate(headers)
                             if value in DISCIPLINE_HEADERS), -1)
    if class_index < 0:
        class_index = next((index for index in range(len(columns))
            if index != enrollment_index and sum(bool(re.fullmatch(r"\d{4,8}",
                row[index].strip())) for row in rows[:30] if index < len(row))
                >= max(2, len(rows[:30]) // 2)), -1)
    if discipline_index < 0:
        discipline_index = next((index for index in range(len(columns))
            if sum(bool(re.search(r"\bARA\d{4}\b", row[index], re.I))
                   for row in rows[:30] if index < len(row))
                >= max(2, len(rows[:30]) // 2)), -1)
    class_values = sorted({row[class_index].strip() for row in rows
                           if class_index >= 0 and class_index < len(row)
                           and row[class_index].strip()})[:20]
    discipline_values = sorted({row[discipline_index].strip() for row in rows
                                if discipline_index >= 0 and discipline_index < len(row)
                                and row[discipline_index].strip()})[:20]
    return dict(class_column=class_index, discipline_column=discipline_index,
                class_values=class_values, discipline_values=discipline_values)


def _preview(source_type: str, rows: list[list[str]], source_text: str = "") -> dict:
    columns, records, enrollment, name, warnings = _mapping(rows)
    return dict(source_type=source_type, columns=columns, rows=records,
                enrollment_column=enrollment, name_column=name,
                warnings=warnings, source_text=source_text,
                **_scope_hints(columns, records, enrollment))


def preview_student_roster_source(data: bytes | None = None, filename: str = "",
                                  pasted_text: str = "") -> dict:
    if data is not None and len(data) > MAX_SOURCE_BYTES:
        raise ValueError("Arquivo maior que 10 MB")
    if len(pasted_text) > 300_000:
        raise ValueError("Texto colado maior que 300 mil caracteres")
    extension = Path(filename).suffix.lower()
    source_type = extension.lstrip(".") or "texto colado"
    if data is not None and extension in IMAGE_EXTENSIONS:
        from . import ocr_service
        if not ocr_service.is_available():
            raise ValueError("OCR não está disponível neste servidor")
        text = ocr_service.text_from_image(data)
        if not text.strip():
            raise ValueError("Não consegui reconhecer texto no print")
        source_type = "print (OCR)"
    elif data is not None and extension == ".xlsx":
        rows, sheet = _xlsx_rows(data)
        source_type = f"XLSX • {sheet}"
        return _preview(source_type, rows)
    else:
        text = _decode(data) if data is not None else pasted_text
    if not text.strip():
        raise ValueError("Fonte vazia")
    if extension == ".xml" or text.lstrip().startswith("<?xml"):
        rows = _xml_rows(text)
        source_type = "XML"
    elif extension == ".json" or text.lstrip().startswith(("[", "{")):
        rows = _json_rows(text)
        source_type = "JSON"
    elif "<table" in text.lower():
        parser = _TableHTMLParser()
        parser.feed(text)
        rows = parser.rows
        source_type = "tabela HTML"
    else:
        rows = _delimited_rows(text)
        if not rows or max(map(len, rows)) < 2:
            pairs = _line_pair_rows(text)
            rows = [["Matrícula", "Nome"], *pairs] if pairs else []
            if source_type != "print (OCR)":
                source_type = "texto/print"
    if not rows:
        raise ValueError("Não encontrei uma tabela ou linhas com matrícula e nome")
    return _preview(source_type, rows,
        text[:1200] if source_type == "print (OCR)" else "")
